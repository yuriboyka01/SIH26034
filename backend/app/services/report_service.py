"""
Phase 5: Report generation service.

Generates PDF and DOCX reports from Phase 4 compliance data.
"""

import io
from datetime import datetime
from uuid import UUID

from sqlalchemy.orm import Session
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import inch
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from app.core.exceptions import NotFoundError, BadRequestError
from app.repositories.inspection_repository import InspectionRepository
from app.repositories.compliance_repository import ComplianceRepository
from app.schemas.report import (
    InspectionReportData,
    ReportProductInfo,
    ReportImageResult,
    ReportRuleResult
)

# Status ordering as per requirements: FAIL > REVIEW > PASS > NOT_APPLICABLE
STATUS_PRIORITY = {
    "FAIL": 4,
    "REVIEW": 3,
    "PASS": 2,
    "NOT_APPLICABLE": 1,
    "NOT_ANALYSED": 0
}


class ReportService:
    def __init__(self, db: Session):
        self.db = db
        self.insp_repo = InspectionRepository(db)
        self.comp_repo = ComplianceRepository(db)

    def generate_report_data(self, inspection_id: UUID, user_id: UUID) -> InspectionReportData:
        """Assembles the full compliance report data from the database."""
        inspection = self.insp_repo.get_by_id(inspection_id)
        if not inspection or inspection.created_by != user_id:
            raise NotFoundError(code="INSPECTION_NOT_FOUND", message="Inspection not found.")

        reports = self.comp_repo.get_reports_by_inspection_id(inspection_id)
        
        overall_inspection_status = "NOT_ANALYSED"
        
        total_rules = 0
        passed = 0
        failed = 0
        review = 0
        not_applicable = 0
        
        image_results = []
        violations = []
        
        # Determine engine and ruleset versions based on the first report, or default if empty
        engine_version = "1.0.0"
        ruleset_version = "LM-2011-v1"
        if reports:
            engine_version = reports[0].engine_version or "1.0.0"
            ruleset_version = reports[0].ruleset_version or "LM-2011-v1"
            
        product_info_data = ReportProductInfo()
        
        for rep in reports:
            # Aggregate status logic: FAIL > REVIEW > PASS > NOT_APPLICABLE
            if STATUS_PRIORITY.get(rep.overall_status, 0) > STATUS_PRIORITY.get(overall_inspection_status, 0):
                overall_inspection_status = rep.overall_status
                
            total_rules += rep.total_rules_checked
            passed += rep.passed_count
            failed += rep.failed_count
            review += rep.review_count
            not_applicable += rep.not_applicable_count
            
            # Map rule results
            rule_results_mapped = []
            for rr in rep.rule_results:
                mapped_rr = ReportRuleResult(
                    rule_id=rr.rule_id,
                    rule_name=rr.rule_name,
                    status=rr.status,
                    severity=rr.severity,
                    message=rr.message,
                    field=rr.field,
                    expected=rr.expected,
                    actual=rr.actual,
                    source_reference=rr.source_reference,
                    evidence=rr.get_evidence()
                )
                rule_results_mapped.append(mapped_rr)
                if rr.status == "FAIL":
                    violations.append(mapped_rr)
                    
            image_results.append(
                ReportImageResult(
                    image_id=str(rep.ocr_result.image_id),
                    overall_status=rep.overall_status,
                    total_rules_checked=rep.total_rules_checked,
                    passed_count=rep.passed_count,
                    failed_count=rep.failed_count,
                    review_count=rep.review_count,
                    not_applicable_count=rep.not_applicable_count,
                    rule_results=rule_results_mapped
                )
            )
            
            # Populate product info if available
            if rep.ocr_result and rep.ocr_result.product_info:
                pinfo = rep.ocr_result.product_info
                # Only take first non-null values if multiple images
                product_info_data.product_name = product_info_data.product_name or pinfo.product_name
                product_info_data.brand_name = product_info_data.brand_name or pinfo.brand_name
                product_info_data.manufacturer = product_info_data.manufacturer or pinfo.manufacturer
                product_info_data.net_quantity = product_info_data.net_quantity or pinfo.net_quantity
                product_info_data.mrp = product_info_data.mrp or pinfo.mrp
                product_info_data.manufacturing_date = product_info_data.manufacturing_date or pinfo.manufacturing_date
                product_info_data.expiry_date = product_info_data.expiry_date or pinfo.expiry_date
                product_info_data.batch_number = product_info_data.batch_number or pinfo.batch_number
                product_info_data.country_of_origin = product_info_data.country_of_origin or pinfo.country_of_origin
                product_info_data.ingredients = product_info_data.ingredients or pinfo.ingredients
                product_info_data.license_number = product_info_data.license_number or pinfo.license_number
                product_info_data.customer_care = product_info_data.customer_care or pinfo.customer_care

        return InspectionReportData(
            report_id=f"REP-{inspection.inspection_number}",
            inspection_id=str(inspection_id),
            inspection_number=inspection.inspection_number,
            inspection_date=inspection.created_at.strftime("%Y-%m-%d %H:%M:%S UTC"),
            product_name=inspection.product_name,
            brand=inspection.brand,
            overall_status=overall_inspection_status,
            product_info=product_info_data,
            image_results=image_results,
            total_rules_checked=total_rules,
            passed_count=passed,
            failed_count=failed,
            review_count=review,
            not_applicable_count=not_applicable,
            violations=violations,
            generated_at=datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC"),
            engine_version=engine_version,
            ruleset_version=ruleset_version
        )

    def render_pdf(self, report_data: InspectionReportData) -> bytes:
        """Generates a PDF report using reportlab."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer, 
            pagesize=A4,
            rightMargin=inch, 
            leftMargin=inch, 
            topMargin=inch, 
            bottomMargin=inch
        )
        
        styles = getSampleStyleSheet()
        title_style = styles['Title']
        h2_style = styles['Heading2']
        h3_style = styles['Heading3']
        normal_style = styles['Normal']
        
        bold_style = ParagraphStyle(
            'BoldStyle',
            parent=normal_style,
            fontName='Helvetica-Bold'
        )
        
        elements = []
        
        # Header
        elements.append(Paragraph("SIH26034 COMPLIANCE INSPECTION REPORT", title_style))
        elements.append(Spacer(1, 0.25 * inch))
        
        # Inspection Info
        elements.append(Paragraph("Inspection Information", h2_style))
        info_data = [
            ["Inspection ID:", report_data.inspection_number],
            ["Date:", report_data.inspection_date],
            ["Product:", f"{report_data.product_name} ({report_data.brand})"],
            ["Overall Status:", report_data.overall_status]
        ]
        info_table = Table(info_data, colWidths=[1.5*inch, 4.5*inch])
        info_table.setStyle(TableStyle([
            ('FONTNAME', (0,0), (0,-1), 'Helvetica-Bold'),
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ]))
        elements.append(info_table)
        elements.append(Spacer(1, 0.25 * inch))
        
        # Summary
        elements.append(Paragraph("SUMMARY", h2_style))
        summary_data = [
            ["Rules Checked:", str(report_data.total_rules_checked)],
            ["Passed:", str(report_data.passed_count)],
            ["Failed:", str(report_data.failed_count)],
            ["Review:", str(report_data.review_count)],
            ["Not Applicable:", str(report_data.not_applicable_count)]
        ]
        summary_table = Table(summary_data, colWidths=[1.5*inch, 4.5*inch])
        summary_table.setStyle(TableStyle([
            ('FONTNAME', (0,0), (0,-1), 'Helvetica-Bold'),
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ]))
        elements.append(summary_table)
        elements.append(Spacer(1, 0.25 * inch))
        
        # Product Information
        elements.append(Paragraph("PRODUCT INFORMATION", h2_style))
        pinfo = report_data.product_info
        if pinfo:
            p_data = []
            if pinfo.product_name: p_data.append(["Product Name:", pinfo.product_name])
            if pinfo.net_quantity: p_data.append(["Net Quantity:", pinfo.net_quantity])
            if pinfo.mrp: p_data.append(["MRP:", pinfo.mrp])
            if pinfo.manufacturer: p_data.append(["Manufacturer:", pinfo.manufacturer])
            if pinfo.manufacturing_date: p_data.append(["Manufacturing Date:", pinfo.manufacturing_date])
            if pinfo.customer_care: p_data.append(["Consumer Care:", pinfo.customer_care])
            
            if p_data:
                # Need Paragraphs to wrap long texts
                wrapped_p_data = [[Paragraph(row[0], bold_style), Paragraph(row[1], normal_style)] for row in p_data]
                p_table = Table(wrapped_p_data, colWidths=[1.5*inch, 4.5*inch])
                p_table.setStyle(TableStyle([
                    ('VALIGN', (0,0), (-1,-1), 'TOP'),
                    ('BOTTOMPADDING', (0,0), (-1,-1), 6),
                ]))
                elements.append(p_table)
            else:
                elements.append(Paragraph("No product information extracted.", normal_style))
        else:
            elements.append(Paragraph("No product information available.", normal_style))
            
        elements.append(Spacer(1, 0.25 * inch))
        
        # Violations
        elements.append(Paragraph("VIOLATIONS", h2_style))
        if report_data.violations:
            for viol in report_data.violations:
                elements.append(Paragraph(f"{viol.rule_id}: {viol.rule_name}", h3_style))
                v_data = [
                    ["Status:", viol.status],
                    ["Severity:", viol.severity],
                    ["Message:", viol.message]
                ]
                wrapped_v = [[Paragraph(row[0], bold_style), Paragraph(str(row[1]), normal_style)] for row in v_data]
                v_table = Table(wrapped_v, colWidths=[1*inch, 5*inch])
                v_table.setStyle(TableStyle([('VALIGN', (0,0), (-1,-1), 'TOP')]))
                elements.append(v_table)
                elements.append(Spacer(1, 0.1 * inch))
        else:
            elements.append(Paragraph("No violations found.", normal_style))
            
        elements.append(Spacer(1, 0.25 * inch))
        
        # Rule Results Details
        elements.append(Paragraph("RULE RESULTS DETAILS", h2_style))
        for img_res in report_data.image_results:
            for rr in img_res.rule_results:
                elements.append(Paragraph(f"{rr.rule_id} - {rr.rule_name}", h3_style))
                
                rr_data = [
                    ["Status:", rr.status],
                    ["Expected:", rr.expected],
                    ["Actual:", rr.actual if rr.actual else "Not detected"],
                ]
                
                wrapped_rr = [[Paragraph(row[0], bold_style), Paragraph(str(row[1]), normal_style)] for row in rr_data]
                rr_table = Table(wrapped_rr, colWidths=[1*inch, 5*inch])
                rr_table.setStyle(TableStyle([('VALIGN', (0,0), (-1,-1), 'TOP'), ('BOTTOMPADDING', (0,0), (-1,-1), 4)]))
                elements.append(rr_table)
                
                elements.append(Spacer(1, 0.05 * inch))
                elements.append(Paragraph("EVIDENCE", bold_style))
                
                if rr.evidence and rr.evidence.get("source_text"):
                    evidence_text = f"Source Text: {rr.evidence.get('source_text')} (Confidence: {rr.evidence.get('confidence', 0):.2f})"
                    if rr.evidence.get("bbox"):
                        evidence_text += f"\nBBox: {rr.evidence.get('bbox')}"
                    elements.append(Paragraph(evidence_text.replace("\n", "<br/>"), normal_style))
                else:
                    elements.append(Paragraph("No evidence available.", normal_style))
                    
                elements.append(Spacer(1, 0.15 * inch))
                
        # Metadata
        elements.append(Spacer(1, 0.5 * inch))
        elements.append(Paragraph("REPORT METADATA", h2_style))
        meta_data = [
            ["Generated At:", report_data.generated_at],
            ["Ruleset Version:", report_data.ruleset_version or "N/A"],
            ["Engine Version:", report_data.engine_version or "N/A"]
        ]
        meta_table = Table(meta_data, colWidths=[1.5*inch, 4.5*inch])
        meta_table.setStyle(TableStyle([
            ('FONTNAME', (0,0), (0,-1), 'Helvetica-Bold'),
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ]))
        elements.append(meta_table)

        doc.build(elements)
        return buffer.getvalue()

    def render_docx(self, report_data: InspectionReportData) -> bytes:
        """Generates a DOCX report using python-docx."""
        try:
            import docx
            from docx.shared import Inches, Pt
            from docx.enum.text import WD_ALIGN_PARAGRAPH
        except Exception as exc:
            raise BadRequestError(
                code="DOCX_UNAVAILABLE",
                message="DOCX generation library (python-docx) is not installed or available."
            ) from exc

        doc = docx.Document()
        
        # Title
        title = doc.add_heading('SIH26034 COMPLIANCE INSPECTION REPORT', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Inspection Info
        doc.add_heading('Inspection Information', level=1)
        doc.add_paragraph(f"Inspection ID: {report_data.inspection_number}")
        doc.add_paragraph(f"Date: {report_data.inspection_date}")
        doc.add_paragraph(f"Product: {report_data.product_name} ({report_data.brand})")
        doc.add_paragraph(f"Overall Status: {report_data.overall_status}")
        
        # Summary
        doc.add_heading('SUMMARY', level=1)
        doc.add_paragraph(f"Rules Checked: {report_data.total_rules_checked}")
        doc.add_paragraph(f"Passed: {report_data.passed_count}")
        doc.add_paragraph(f"Failed: {report_data.failed_count}")
        doc.add_paragraph(f"Review: {report_data.review_count}")
        doc.add_paragraph(f"Not Applicable: {report_data.not_applicable_count}")
        
        # Product Information
        doc.add_heading('PRODUCT INFORMATION', level=1)
        pinfo = report_data.product_info
        if pinfo:
            if pinfo.product_name: doc.add_paragraph(f"Product Name: {pinfo.product_name}")
            if pinfo.net_quantity: doc.add_paragraph(f"Net Quantity: {pinfo.net_quantity}")
            if pinfo.mrp: doc.add_paragraph(f"MRP: {pinfo.mrp}")
            if pinfo.manufacturer: doc.add_paragraph(f"Manufacturer: {pinfo.manufacturer}")
            if pinfo.manufacturing_date: doc.add_paragraph(f"Manufacturing Date: {pinfo.manufacturing_date}")
            if pinfo.customer_care: doc.add_paragraph(f"Consumer Care: {pinfo.customer_care}")
        else:
            doc.add_paragraph("No product information available.")
            
        # Violations
        doc.add_heading('VIOLATIONS', level=1)
        if report_data.violations:
            for viol in report_data.violations:
                doc.add_heading(f"{viol.rule_id}: {viol.rule_name}", level=2)
                doc.add_paragraph(f"Status: {viol.status}")
                doc.add_paragraph(f"Severity: {viol.severity}")
                doc.add_paragraph(f"Message: {viol.message}")
        else:
            doc.add_paragraph("No violations found.")
            
        # Rule Results Details
        doc.add_heading('RULE RESULTS DETAILS', level=1)
        for img_res in report_data.image_results:
            for rr in img_res.rule_results:
                doc.add_heading(f"{rr.rule_id} - {rr.rule_name}", level=2)
                doc.add_paragraph(f"Status: {rr.status}")
                doc.add_paragraph(f"Expected: {rr.expected}")
                doc.add_paragraph(f"Actual: {rr.actual if rr.actual else 'Not detected'}")
                
                # Evidence
                evidence_p = doc.add_paragraph()
                evidence_p.add_run('EVIDENCE\n').bold = True
                
                if rr.evidence and rr.evidence.get("source_text"):
                    evidence_p.add_run(f"Source Text: {rr.evidence.get('source_text')} (Confidence: {rr.evidence.get('confidence', 0):.2f})\n")
                    if rr.evidence.get("bbox"):
                        evidence_p.add_run(f"BBox: {rr.evidence.get('bbox')}")
                else:
                    evidence_p.add_run("No evidence available.")
                    
        # Metadata
        doc.add_heading('REPORT METADATA', level=1)
        doc.add_paragraph(f"Generated At: {report_data.generated_at}")
        doc.add_paragraph(f"Ruleset Version: {report_data.ruleset_version or 'N/A'}")
        doc.add_paragraph(f"Engine Version: {report_data.engine_version or 'N/A'}")
        
        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    def render_show_cause_pdf(self, report_data: InspectionReportData) -> bytes:
        """Generates an Auto-Drafted Show-Cause Notice PDF."""
        from app.core.config import settings
        
        inspection = self.insp_repo.get_by_id(UUID(report_data.inspection_id))
        if not inspection:
            raise NotFoundError(code="INSPECTION_NOT_FOUND", message="Inspection not found.")

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer, 
            pagesize=A4,
            rightMargin=inch, 
            leftMargin=inch, 
            topMargin=inch, 
            bottomMargin=inch
        )
        
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'NoticeTitle',
            parent=styles['Heading1'],
            alignment=1, # Center
            spaceAfter=12
        )
        h2_style = styles['Heading2']
        h3_style = styles['Heading3']
        normal_style = styles['Normal']
        bold_style = ParagraphStyle('BoldStyle', parent=normal_style, fontName='Helvetica-Bold')
        warning_style = ParagraphStyle('WarningStyle', parent=normal_style, textColor=colors.red, alignment=1, spaceAfter=12)
        
        elements = []
        
        # Draft Warning
        elements.append(Paragraph("DRAFT — FOR REVIEW AND AUTHORIZATION", warning_style))
        
        # Header
        elements.append(Paragraph("NOTICE TO SHOW CAUSE", title_style))
        elements.append(Paragraph("under Section 36 & 48 of Legal Metrology Act, 2009", ParagraphStyle('SubTitle', parent=normal_style, alignment=1, spaceAfter=24)))
        
        # Metadata
        current_date = datetime.now().strftime("%Y-%m-%d")
        meta_data = [
            ["Notice Reference Number:", f"SCN-{report_data.inspection_number}"],
            ["Notice Date:", current_date],
            ["Inspection Number:", report_data.inspection_number],
            ["Inspection Date:", report_data.inspection_date],
        ]
        meta_table = Table(meta_data, colWidths=[2.5*inch, 3.5*inch])
        meta_table.setStyle(TableStyle([
            ('FONTNAME', (0,0), (0,-1), 'Helvetica-Bold'),
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ]))
        elements.append(meta_table)
        elements.append(Spacer(1, 0.25 * inch))
        
        # Addressee
        elements.append(Paragraph("To:", bold_style))
        establishment_name = inspection.establishment_name or "The Operator / Manufacturer / Packer"
        elements.append(Paragraph(establishment_name, normal_style))
        if inspection.latitude and inspection.longitude:
            elements.append(Paragraph(f"GPS Location: {inspection.latitude}, {inspection.longitude}", normal_style))
        else:
            elements.append(Paragraph("GPS Location: Not recorded", normal_style))
            
        pinfo = report_data.product_info
        manufacturer = (pinfo.manufacturer if pinfo else None) or report_data.brand or "Not recorded"
        elements.append(Paragraph(f"Manufacturer / Brand: {manufacturer}", normal_style))
        elements.append(Spacer(1, 0.25 * inch))
        
        # Subject
        elements.append(Paragraph("Subject: Observations regarding non-compliance with Packaged Commodities Rules, 2011", bold_style))
        elements.append(Spacer(1, 0.25 * inch))
        
        # Inspection Statement
        elements.append(Paragraph(f"An inspection was conducted on {report_data.inspection_date}. The inspection records indicate the following observed non-compliances concerning the product '{report_data.product_name}':", normal_style))
        elements.append(Spacer(1, 0.15 * inch))
        
        # Contraventions Table
        violations = [v for v in report_data.violations if v.status == "FAIL"]
        if violations:
            table_data = [["Rule Reference", "Expected Declaration", "Actual Observation"]]
            for v in violations:
                rule_ref = Paragraph(f"{v.rule_name}<br/>({v.source_reference})", normal_style)
                expected = Paragraph(str(v.expected), normal_style)
                actual = Paragraph(str(v.actual) if v.actual else "Not established", normal_style)
                table_data.append([rule_ref, expected, actual])
                
            v_table = Table(table_data, colWidths=[2.0*inch, 2.0*inch, 2.0*inch])
            v_table.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.lightgrey),
                ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
                ('VALIGN', (0,0), (-1,-1), 'TOP'),
                ('GRID', (0,0), (-1,-1), 1, colors.black),
                ('BOTTOMPADDING', (0,0), (-1,-1), 6),
            ]))
            elements.append(v_table)
        else:
            elements.append(Paragraph("No failed rules identified in the compliance report.", normal_style))
            
        elements.append(Spacer(1, 0.25 * inch))
        
        # Warning / Deadline
        deadline_days = settings.SHOW_CAUSE_RESPONSE_DAYS
        elements.append(Paragraph(f"You are requested to show cause within {deadline_days} days from receipt of this notice as to why appropriate action should not be taken under the applicable provisions of the Legal Metrology Act, 2009 and the rules made thereunder.", normal_style))
        elements.append(Spacer(1, 0.5 * inch))
        
        # Signature block
        elements.append(Paragraph("Authorized Signatory", normal_style))
        elements.append(Paragraph("_______________________", normal_style))
        elements.append(Paragraph("Legal Metrology Officer", normal_style))

        doc.build(elements)
        return buffer.getvalue()

    def render_show_cause_docx(self, report_data: InspectionReportData) -> bytes:
        """Generates an Auto-Drafted Show-Cause Notice DOCX."""
        try:
            import docx
            from docx.shared import Inches, Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH
        except Exception as exc:
            raise BadRequestError(
                code="DOCX_UNAVAILABLE",
                message="DOCX generation library (python-docx) is not installed or available."
            ) from exc

        from app.core.config import settings
        
        inspection = self.insp_repo.get_by_id(UUID(report_data.inspection_id))
        if not inspection:
            raise NotFoundError(code="INSPECTION_NOT_FOUND", message="Inspection not found.")

        doc = docx.Document()
        
        # Draft Warning
        warning_p = doc.add_paragraph()
        warning_run = warning_p.add_run("DRAFT — FOR REVIEW AND AUTHORIZATION")
        warning_run.bold = True
        warning_run.font.color.rgb = RGBColor(255, 0, 0)
        warning_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Header
        title = doc.add_heading('NOTICE TO SHOW CAUSE', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        subtitle = doc.add_paragraph('under Section 36 & 48 of Legal Metrology Act, 2009')
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Metadata
        current_date = datetime.now().strftime("%Y-%m-%d")
        doc.add_paragraph(f"Notice Reference Number: SCN-{report_data.inspection_number}").bold = True
        doc.add_paragraph(f"Notice Date: {current_date}")
        doc.add_paragraph(f"Inspection Number: {report_data.inspection_number}")
        doc.add_paragraph(f"Inspection Date: {report_data.inspection_date}")
        
        doc.add_paragraph() # Spacer
        
        # Addressee
        doc.add_paragraph("To:").bold = True
        establishment_name = inspection.establishment_name or "The Operator / Manufacturer / Packer"
        doc.add_paragraph(establishment_name)
        if inspection.latitude and inspection.longitude:
            doc.add_paragraph(f"GPS Location: {inspection.latitude}, {inspection.longitude}")
        else:
            doc.add_paragraph("GPS Location: Not recorded")
            
        pinfo = report_data.product_info
        manufacturer = (pinfo.manufacturer if pinfo else None) or report_data.brand or "Not recorded"
        doc.add_paragraph(f"Manufacturer / Brand: {manufacturer}")
        
        doc.add_paragraph() # Spacer
        
        # Subject
        doc.add_paragraph("Subject: Observations regarding non-compliance with Packaged Commodities Rules, 2011").bold = True
        doc.add_paragraph()
        
        # Body
        doc.add_paragraph(f"An inspection was conducted on {report_data.inspection_date}. The inspection records indicate the following observed non-compliances concerning the product '{report_data.product_name}':")
        
        # Table
        violations = [v for v in report_data.violations if v.status == "FAIL"]
        if violations:
            table = doc.add_table(rows=1, cols=3)
            table.style = 'Table Grid'
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = 'Rule Reference'
            hdr_cells[1].text = 'Expected Declaration'
            hdr_cells[2].text = 'Actual Observation'
            
            # Bold headers
            for cell in hdr_cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True

            for v in violations:
                row_cells = table.add_row().cells
                row_cells[0].text = f"{v.rule_name}\n({v.source_reference})"
                row_cells[1].text = str(v.expected)
                row_cells[2].text = str(v.actual) if v.actual else "Not established"
        else:
            doc.add_paragraph("No failed rules identified in the compliance report.")
            
        doc.add_paragraph() # Spacer
        
        # Deadline
        deadline_days = settings.SHOW_CAUSE_RESPONSE_DAYS
        doc.add_paragraph(f"You are requested to show cause within {deadline_days} days from receipt of this notice as to why appropriate action should not be taken under the applicable provisions of the Legal Metrology Act, 2009 and the rules made thereunder.")
        
        # Signature
        doc.add_paragraph("\n\nAuthorized Signatory")
        doc.add_paragraph("_______________________")
        doc.add_paragraph("Legal Metrology Officer")
        
        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()
